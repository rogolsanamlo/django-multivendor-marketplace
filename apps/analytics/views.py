import json

import pandas as pd
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.views import View
from openpyxl.styles import Font

from .pandas_services import (
    abc_analysis,
    customer_segmentation,
    monthly_trend_dataframe,
    sales_by_product,
    sales_dataframe,
    sales_forecast,
    sales_growth_analysis,
    sales_pivot_table,
    weekday_sales_analysis,
)
from .services import (
    get_active_products_count,
    get_average_order_amount,
    get_available_products_count,
    get_daily_sales,
    get_last_7_days_sales,
    get_last_orders,
    get_low_stock_products,
    get_monthly_income,
    get_new_customers_this_month,
    get_orders_count,
    get_products_without_sales,
    get_store_profit,
    get_top_categories,
    get_top_customers,
    get_top_products,
)


class StoreRequiredMixin(LoginRequiredMixin):
    """Logged-in users without a store are sent to the store creation page."""

    def dispatch(self, request, *args, **kwargs):
        if (
            request.user.is_authenticated
            and not hasattr(request.user, "store")
        ):
            messages.error(
                request,
                "برای مشاهده گزارش‌ها ابتدا فروشگاه خود را ثبت کنید.",
            )

            return redirect("accounts:create_store")

        return super().dispatch(request, *args, **kwargs)


class SellerAnalyticsDashboardView(
    StoreRequiredMixin,
    View,
):

    def get(self, request):

        store = request.user.store

        daily_sales = get_daily_sales(
            store
        )

        daily_labels = []
        daily_values = []

        for item in daily_sales:
            daily_labels.append(
                item[
                    "order__register_date"
                ].strftime("%Y-%m-%d")
            )

            daily_values.append(
                item["total_sales"]
            )

        monthly_income = get_monthly_income(
            store
        )

        month_labels = []
        month_values = []

        for item in monthly_income:
            month_labels.append(
                item["month"].strftime("%Y-%m")
            )

            month_values.append(
                item["income"]
            )

        last_7_days = get_last_7_days_sales(
            store
        )

        week_labels = []
        week_values = []

        for item in last_7_days:
            week_labels.append(
                item["date"].strftime("%m/%d")
            )

            week_values.append(
                item["amount"]
            )

        top_products = get_top_products(
            store
        )

        top_categories = get_top_categories(
            store
        )

        group_key = (
            "store_product__product__product_group__group_title"
        )

        group_labels = [
            item[group_key]
            for item in top_categories
        ]

        group_values = [
            item["total_sales"]
            for item in top_categories
        ]

        products_without_sales = (
            get_products_without_sales(store)
        )

        low_stock_products = (
            get_low_stock_products(store)
        )

        orders_count = get_orders_count(
            store
        )

        active_products_count = (
            get_active_products_count(store)
        )

        available_products_count = (
            get_available_products_count(store)
        )

        profit = get_store_profit(
            store
        )

        last_orders = get_last_orders(
            store
        )

        top_customers = get_top_customers(
            store
        )

        new_customers = (
            get_new_customers_this_month(store)
        )

        average_order = (
            get_average_order_amount(store)
        )

        context = {
            "top_products": top_products,
            "top_categories": top_categories,
            "group_labels": json.dumps(
                group_labels
            ),
            "group_values": json.dumps(
                group_values
            ),
            "daily_sales": daily_sales,
            "daily_labels": json.dumps(
                daily_labels
            ),
            "daily_values": json.dumps(
                daily_values
            ),
            "monthly_income": monthly_income,
            "month_labels": json.dumps(
                month_labels
            ),
            "month_values": json.dumps(
                month_values
            ),
            "products_without_sales": (
                products_without_sales
            ),
            "low_stock_products": (
                low_stock_products
            ),
            "orders_count": orders_count,
            "active_products_count": (
                active_products_count
            ),
            "available_products_count": (
                available_products_count
            ),
            "profit": profit,
            "last_orders": last_orders,
            "top_customers": top_customers,
            "week_labels": json.dumps(
                week_labels
            ),
            "week_values": json.dumps(
                week_values
            ),
            "new_customers": new_customers,
            "average_order": average_order,
        }

        return render(
            request,
            "analytics/dashboard.html",
            context,
        )


class ReportsView(
    StoreRequiredMixin,
    View,
):

    def get(self, request):

        store = request.user.store

        pandas_summary = sales_dataframe(
            store
        )

        sales_products = sales_by_product(
            store
        )

        pivot_table = sales_pivot_table(
            store
        )

        name_key = "store_product__product__product_name"

        pivot_months = sorted({
            str(key)
            for record in pivot_table
            for key in record
            if key != name_key
        })

        pivot_rows = [
            {
                "name": record[name_key],
                "cells": [
                    record.get(month_key, 0)
                    for month_key in sorted(
                        (k for k in record if k != name_key),
                        key=str,
                    )
                ],
            }
            for record in pivot_table
        ]

        monthly_trend = (
            monthly_trend_dataframe(store)
        )

        month_labels = [
            item["month"]
            for item in monthly_trend
        ]

        month_values = [
            item["total_price"]
            for item in monthly_trend
        ]

        abc_result = abc_analysis(
            store
        )

        abc_labels = [
            item[
                "store_product__product__product_name"
            ]
            for item in abc_result
        ]

        abc_values = [
            item["total_price"]
            for item in abc_result
        ]

        customers = customer_segmentation(
            store
        )

        weekday_sales = (
            weekday_sales_analysis(store)
        )

        weekday_labels = [
            item["weekday"]
            for item in weekday_sales
        ]

        weekday_values = [
            item["total_price"]
            for item in weekday_sales
        ]

        sales_growth = (
            sales_growth_analysis(store)
        )

        growth_labels = [
            item["month"]
            for item in sales_growth
        ]

        growth_values = [
            item["growth_percent"]
            for item in sales_growth
        ]

        forecast = sales_forecast(
            store
        )

        context = {
            "pandas_summary": pandas_summary,
            "sales_products": (
                sales_products.to_dict(
                    "records"
                )
                if not sales_products.empty
                else []
            ),
            "pivot_table": pivot_table,
            "pivot_months": pivot_months,
            "pivot_rows": pivot_rows,
            "monthly_trend": monthly_trend,
            "abc_analysis": abc_result,
            "customer_segmentation": customers,
            "weekday_sales": weekday_sales,
            "sales_growth": sales_growth,
            "month_labels": json.dumps(
                month_labels
            ),
            "month_values": json.dumps(
                month_values
            ),
            "weekday_labels": json.dumps(
                weekday_labels
            ),
            "weekday_values": json.dumps(
                weekday_values
            ),
            "growth_labels": json.dumps(
                growth_labels
            ),
            "growth_values": json.dumps(
                growth_values
            ),
            "abc_labels": json.dumps(
                abc_labels
            ),
            "abc_values": json.dumps(
                abc_values
            ),
            "forecast": forecast,
        }

        return render(
            request,
            "analytics/reports.html",
            context,
        )


class ExportSalesExcelView(
    StoreRequiredMixin,
    View,
):

    def get(self, request):

        store = request.user.store

        df = sales_by_product(
            store
        )

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )

        response["Content-Disposition"] = (
            'attachment; filename="sales_report.xlsx"'
        )

        report = df.rename(
            columns={
                "store_product__product__product_name": "محصول",
                "total_price": "مبلغ فروش (تومان)",
            }
        ) if not df.empty else pd.DataFrame(
            columns=["محصول", "مبلغ فروش (تومان)"]
        )

        sheet_name = "گزارش فروش"

        with pd.ExcelWriter(
            response,
            engine="openpyxl",
        ) as writer:

            report.to_excel(
                writer,
                index=False,
                sheet_name=sheet_name,
            )

            sheet = writer.sheets[sheet_name]
            sheet.sheet_view.rightToLeft = True
            sheet.column_dimensions["A"].width = 40
            sheet.column_dimensions["B"].width = 24

            for cell in sheet[1]:
                cell.font = Font(bold=True)

            last_row = len(report) + 1

            for row in range(2, last_row + 1):
                sheet.cell(row=row, column=2).number_format = "#,##0"

            if len(report):
                total_row = last_row + 1
                sheet.cell(row=total_row, column=1, value="جمع کل")
                total_cell = sheet.cell(
                    row=total_row,
                    column=2,
                    value=f"=SUM(B2:B{last_row})",
                )
                total_cell.number_format = "#,##0"
                sheet.cell(row=total_row, column=1).font = Font(bold=True)
                total_cell.font = Font(bold=True)

        return response


class ExportPDFView(
    LoginRequiredMixin,
    View,
):

    def get(self, request):

        messages.warning(
            request,
            "امکان دانلود PDF در نسخه فعلی غیرفعال است."
        )

        return redirect(
            "analytics:reports"
        )


# PDF export is disabled because WeasyPrint
# requires native system libraries that are
# not installed in the current environment.
